from __future__ import annotations

from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches

from mcp_servers.servers.document_intelligence import _pptx_theme as theme
from mcp_servers.servers.document_intelligence._pptx_markdown import add_markdown_bullets, apply_bold_runs

_CHART_TYPES = {
    "bar": XL_CHART_TYPE.COLUMN_CLUSTERED,
    "line": XL_CHART_TYPE.LINE,
    "pie": XL_CHART_TYPE.PIE,
}


class PptxGenerationError(ValueError):
    """Raised when a slide spec is malformed (missing/invalid fields for its type)."""


def _new_presentation() -> Presentation:
    prs = Presentation()
    prs.slide_width = Inches(theme.SLIDE_WIDTH_IN)
    prs.slide_height = Inches(theme.SLIDE_HEIGHT_IN)
    return prs


def _blank_slide(prs: Presentation):
    return prs.slides.add_slide(prs.slide_layouts[6])  # layout 6 = Blank


def _fill_background(slide, color) -> None:
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color


def _add_accent_bar(slide, color) -> None:
    bar = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(0.12), Inches(theme.SLIDE_HEIGHT_IN)
    )
    bar.fill.solid()
    bar.fill.fore_color.rgb = color
    bar.line.fill.background()
    bar.shadow.inherit = False


def _add_textbox(slide, left, top, width, height, text, size, color, bold=False, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    frame = box.text_frame
    frame.word_wrap = True
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    apply_bold_runs(paragraph, text, size, color, bold=bold)
    return box


def _add_footer(slide, slide_number: int, accent) -> None:
    _add_textbox(
        slide, theme.SLIDE_WIDTH_IN - 1.2, theme.SLIDE_HEIGHT_IN - 0.4, 0.9, 0.3,
        str(slide_number), theme.CAPTION_SIZE, theme.MUTED, align=PP_ALIGN.RIGHT,
    )


def _build_title_slide(prs: Presentation, spec: dict[str, Any], accent) -> None:
    slide = _blank_slide(prs)
    _fill_background(slide, theme.LIGHT_BG)
    _add_accent_bar(slide, accent)

    title = spec.get("title", "").strip()
    if not title:
        raise PptxGenerationError("Title slide requires a non-empty 'title'")

    _add_textbox(slide, 0.7, 2.1, theme.SLIDE_WIDTH_IN - 1.4, 1.2, title, theme.TITLE_SIZE, theme.INK, bold=True)

    subtitle = spec.get("subtitle", "")
    if subtitle:
        _add_textbox(slide, 0.7, 3.2, theme.SLIDE_WIDTH_IN - 1.4, 0.6, subtitle, theme.SUBTITLE_SIZE, theme.MUTED)


def _build_section_slide(prs: Presentation, spec: dict[str, Any], accent) -> None:
    slide = _blank_slide(prs)
    _fill_background(slide, accent)

    title = spec.get("title", "").strip()
    if not title:
        raise PptxGenerationError("Section slide requires a non-empty 'title'")

    _add_textbox(
        slide, 0.7, theme.SLIDE_HEIGHT_IN / 2 - 0.6, theme.SLIDE_WIDTH_IN - 1.4, 1.2,
        title, theme.SECTION_TITLE_SIZE, theme.WHITE, bold=True,
    )


def _add_slide_title(slide, title: str, accent) -> None:
    _add_accent_bar(slide, accent)
    _add_textbox(slide, 0.6, 0.35, theme.SLIDE_WIDTH_IN - 1.2, 0.8, title, theme.SLIDE_TITLE_SIZE, theme.INK, bold=True)


def _build_bullets_slide(prs: Presentation, spec: dict[str, Any], accent, slide_number: int) -> None:
    slide = _blank_slide(prs)
    _fill_background(slide, theme.WHITE)

    title = spec.get("title", "").strip()
    if not title:
        raise PptxGenerationError("Bullets slide requires a non-empty 'title'")
    _add_slide_title(slide, title, accent)

    bullets = spec.get("bullets", [])
    if not bullets:
        raise PptxGenerationError("Bullets slide requires a non-empty 'bullets' list")

    box = slide.shapes.add_textbox(Inches(0.8), Inches(1.4), Inches(theme.SLIDE_WIDTH_IN - 1.6), Inches(3.6))
    box.text_frame.word_wrap = True
    add_markdown_bullets(box.text_frame, bullets, theme.BODY_SIZE, theme.INK)

    _add_footer(slide, slide_number, accent)


def _build_table_slide(prs: Presentation, spec: dict[str, Any], accent, slide_number: int) -> None:
    slide = _blank_slide(prs)
    _fill_background(slide, theme.WHITE)

    title = spec.get("title", "").strip()
    if not title:
        raise PptxGenerationError("Table slide requires a non-empty 'title'")
    _add_slide_title(slide, title, accent)

    headers = spec.get("headers", [])
    rows = spec.get("rows", [])
    if not headers or not rows:
        raise PptxGenerationError("Table slide requires non-empty 'headers' and 'rows'")

    n_rows, n_cols = len(rows) + 1, len(headers)
    table_shape = slide.shapes.add_table(
        n_rows, n_cols, Inches(0.8), Inches(1.5), Inches(theme.SLIDE_WIDTH_IN - 1.6), Inches(3.4)
    )
    table = table_shape.table

    for col, header in enumerate(headers):
        cell = table.cell(0, col)
        paragraph = cell.text_frame.paragraphs[0]
        apply_bold_runs(paragraph, header, theme.BODY_SIZE, theme.WHITE, bold=True)
        cell.fill.solid()
        cell.fill.fore_color.rgb = accent

    for r, row in enumerate(rows, start=1):
        for c, value in enumerate(row):
            cell = table.cell(r, c)
            apply_bold_runs(cell.text_frame.paragraphs[0], value, theme.BODY_SIZE, theme.INK)

    _add_footer(slide, slide_number, accent)


def _build_chart_slide(prs: Presentation, spec: dict[str, Any], accent, slide_number: int) -> None:
    slide = _blank_slide(prs)
    _fill_background(slide, theme.WHITE)

    title = spec.get("title", "").strip()
    if not title:
        raise PptxGenerationError("Chart slide requires a non-empty 'title'")
    _add_slide_title(slide, title, accent)

    chart_type = spec.get("chart_type", "bar")
    if chart_type not in _CHART_TYPES:
        raise PptxGenerationError(f"chart_type must be one of {list(_CHART_TYPES)}, got {chart_type!r}")

    categories = spec.get("categories", [])
    series = spec.get("series", {})
    if not categories or not series:
        raise PptxGenerationError("Chart slide requires non-empty 'categories' and 'series'")

    chart_data = CategoryChartData()
    chart_data.categories = categories
    for series_name, values in series.items():
        chart_data.add_series(series_name, values)

    slide.shapes.add_chart(
        _CHART_TYPES[chart_type], Inches(0.8), Inches(1.5), Inches(theme.SLIDE_WIDTH_IN - 1.6), Inches(3.4), chart_data
    )

    _add_footer(slide, slide_number, accent)


def _build_comparison_slide(prs: Presentation, spec: dict[str, Any], accent, slide_number: int) -> None:
    slide = _blank_slide(prs)
    _fill_background(slide, theme.WHITE)

    title = spec.get("title", "").strip()
    if not title:
        raise PptxGenerationError("Comparison slide requires a non-empty 'title'")
    _add_slide_title(slide, title, accent)

    left = spec.get("left", {})
    right = spec.get("right", {})
    if not left.get("heading") or not right.get("heading"):
        raise PptxGenerationError("Comparison slide requires 'left' and 'right', each with a 'heading'")

    column_width = (theme.SLIDE_WIDTH_IN - 2.0) / 2

    for column, x_offset in ((left, 0.7), (right, 0.7 + column_width + 0.6)):
        heading_bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x_offset), Inches(1.4), Inches(column_width), Inches(0.5))
        heading_bg.fill.solid()
        heading_bg.fill.fore_color.rgb = accent
        heading_bg.line.fill.background()
        heading_bg.shadow.inherit = False

        _add_textbox(slide, x_offset + 0.15, 1.47, column_width - 0.3, 0.4, column["heading"], theme.BODY_SIZE, theme.WHITE, bold=True)

        points = column.get("points", [])
        box = slide.shapes.add_textbox(Inches(x_offset), Inches(2.1), Inches(column_width), Inches(2.8))
        frame = box.text_frame
        frame.word_wrap = True
        add_markdown_bullets(frame, points, theme.CAPTION_SIZE, theme.INK)

    _add_footer(slide, slide_number, accent)


_BUILDERS = {
    "title": _build_title_slide,
    "section": _build_section_slide,
    "bullets": _build_bullets_slide,
    "table": _build_table_slide,
    "chart": _build_chart_slide,
    "comparison": _build_comparison_slide,
}

_NO_NUMBER_TYPES = {"title", "section"}


def create_presentation(path: Path, slides: list[dict[str, Any]], accent_color: str | None = None) -> None:
    """Build a real, valid .pptx file from a list of slide specs.

    Each spec is a dict with a 'type' key (title/section/bullets/table/chart/comparison)
    plus type-specific fields. Uses python-pptx's real shape/table/chart/text APIs on
    blank layouts with an explicit theme — never a bare-bones default-template dump.
    """
    if not slides:
        raise PptxGenerationError("slides must be a non-empty list")

    accent = theme.accent_color(accent_color)
    prs = _new_presentation()

    numbered_slide_count = 0
    for spec in slides:
        slide_type = spec.get("type")
        builder = _BUILDERS.get(slide_type)
        if builder is None:
            raise PptxGenerationError(f"Unknown slide type {slide_type!r}; must be one of {list(_BUILDERS)}")

        if slide_type in _NO_NUMBER_TYPES:
            builder(prs, spec, accent)
        else:
            numbered_slide_count += 1
            builder(prs, spec, accent, numbered_slide_count)

    prs.save(str(path))
