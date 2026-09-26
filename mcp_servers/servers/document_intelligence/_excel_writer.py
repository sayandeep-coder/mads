from __future__ import annotations

from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.chart import AreaChart, BarChart, LineChart, PieChart, Reference, ScatterChart
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

_HEADER_FILL = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")
_HEADER_FONT = Font(color="FFFFFF", bold=True)

_CHART_CLASSES: dict[str, type] = {
    "bar": BarChart,
    "column": BarChart,
    "line": LineChart,
    "pie": PieChart,
    "scatter": ScatterChart,
    "area": AreaChart,
}

# Named formats callers can ask for by name instead of a raw Excel number-
# format code — covers the common real-world cases (money, percentages,
# dates) without requiring the model to know Excel's format-code syntax.
_NUMBER_FORMATS: dict[str, str] = {
    "currency": '"₹"#,##0.00',
    "currency_usd": '"$"#,##0.00',
    "percent": "0.0%",
    "integer": "#,##0",
    "decimal": "#,##0.00",
    "date": "yyyy-mm-dd",
}


def _col_letter_for_index(columns: list[str], name_or_index: str | int) -> str:
    """Resolve a column reference (either its header name or a 0-based
    index) to a spreadsheet column letter, so callers can specify a chart
    axis or formatted range by column name rather than counting letters.
    """
    if isinstance(name_or_index, int):
        return get_column_letter(name_or_index + 1)
    if name_or_index in columns:
        return get_column_letter(columns.index(name_or_index) + 1)
    raise ValueError(f"Unknown column {name_or_index!r}; must be one of {columns}")


def _apply_table(sheet, columns: list[str], row_count: int, table_spec: dict[str, Any]) -> None:
    """Turn the header+data range into a real Excel Table (filterable,
    sortable, banded rows) — not just styled cells. table_spec:
    {name?, style?} — style is one of openpyxl's built-in TableStyleInfo
    names (defaults to a clean medium-blue banded style).
    """
    if not columns or row_count == 0:
        return
    last_col = get_column_letter(len(columns))
    last_row = row_count + 1  # +1 for the header row
    table_range = f"A1:{last_col}{last_row}"
    table_name = (table_spec.get("name") or "Table1").replace(" ", "_")

    table = Table(displayName=table_name, ref=table_range)
    table.tableStyleInfo = TableStyleInfo(
        name=table_spec.get("style") or "TableStyleMedium9",
        showRowStripes=True,
        showFirstColumn=False,
        showLastColumn=False,
        showColumnStripes=False,
    )
    sheet.add_table(table)


def _apply_number_formats(sheet, columns: list[str], row_count: int, formats: dict[str, str]) -> None:
    """formats: {column_name_or_index: format_name_or_raw_code}. Applies
    the format to every data cell (not the header) in that column.
    """
    for col_ref, format_spec in formats.items():
        col_letter = _col_letter_for_index(columns, col_ref)
        number_format = _NUMBER_FORMATS.get(format_spec, format_spec)
        for row in range(2, row_count + 2):
            sheet[f"{col_letter}{row}"].number_format = number_format


def _apply_conditional_formats(sheet, columns: list[str], row_count: int, rules: list[dict[str, Any]]) -> None:
    """rules: list of {column, type: 'color_scale'|'highlight', ...}.

    'color_scale': a real 2- or 3-color gradient across the column's
    values (low->mid->high) — the standard "heatmap" conditional format.
    'highlight': highlight cells matching a comparison, e.g.
    {type: 'highlight', column: 'Status', operator: 'equal', value:
    'REGRET', fill: 'FFC7CE', font: '9C0006'} to flag problem rows.
    """
    for rule_spec in rules:
        col_letter = _col_letter_for_index(columns, rule_spec["column"])
        cell_range = f"{col_letter}2:{col_letter}{row_count + 1}"

        if rule_spec.get("type") == "color_scale":
            colors = rule_spec.get("colors") or ["F8696B", "FFEB84", "63BE7B"]
            if len(colors) == 2:
                rule = ColorScaleRule(
                    start_type="min", start_color=colors[0], end_type="max", end_color=colors[1]
                )
            else:
                rule = ColorScaleRule(
                    start_type="min",
                    start_color=colors[0],
                    mid_type="percentile",
                    mid_value=50,
                    mid_color=colors[1],
                    end_type="max",
                    end_color=colors[2],
                )
            sheet.conditional_formatting.add(cell_range, rule)

        elif rule_spec.get("type") == "highlight":
            fill = PatternFill(
                start_color=rule_spec.get("fill", "FFC7CE"),
                end_color=rule_spec.get("fill", "FFC7CE"),
                fill_type="solid",
            )
            font = Font(color=rule_spec.get("font", "9C0006"))
            operator = rule_spec.get("operator", "equal")
            value = rule_spec.get("value")
            formula = [f'"{value}"'] if isinstance(value, str) else [str(value)]
            rule = CellIsRule(operator=operator, formula=formula, fill=fill, font=font)
            sheet.conditional_formatting.add(cell_range, rule)


def _apply_charts(sheet, columns: list[str], row_count: int, charts: list[dict[str, Any]]) -> None:
    """charts: list of {type, title?, categories, series: [str,...],
    anchor?}. `categories` and each entry in `series` are column names
    (or 0-based indices) from the same sheet's data range — real, live
    Excel charts referencing the sheet's own cells, not a static image, so
    they stay correct if the data is edited later.
    """
    for chart_spec in charts:
        chart_type = chart_spec.get("type", "bar")
        chart_class = _CHART_CLASSES.get(chart_type)
        if chart_class is None:
            raise ValueError(f"Unknown chart type {chart_type!r}; must be one of {list(_CHART_CLASSES)}")

        chart = chart_class()
        chart.title = chart_spec.get("title")
        chart.style = 10
        if chart_type in ("bar", "column"):
            chart.type = "col" if chart_type == "column" else "bar"

        last_row = row_count + 1
        cat_col = _col_letter_for_index(columns, chart_spec["categories"])
        cat_col_index = openpyxl.utils.column_index_from_string(cat_col)
        categories = Reference(sheet, min_col=cat_col_index, min_row=2, max_row=last_row)

        for series_ref in chart_spec.get("series", []):
            series_col = _col_letter_for_index(columns, series_ref)
            series_col_index = openpyxl.utils.column_index_from_string(series_col)
            data = Reference(sheet, min_col=series_col_index, min_row=1, max_row=last_row)
            chart.add_data(data, titles_from_data=True)

        chart.set_categories(categories)
        anchor = chart_spec.get("anchor") or get_column_letter(len(columns) + 2) + "2"
        sheet.add_chart(chart, anchor)


def create_excel(path: Path, sheets: list[dict[str, Any]]) -> None:
    """Build a real, valid .xlsx workbook at `path` from a list of sheet
    specs, each:
        {
            name, columns: [str,...], rows: [[...],...],
            table?: {name?, style?},
            number_formats?: {column: 'currency'|'percent'|'date'|... | raw code},
            conditional_formats?: [{column, type: 'color_scale'|'highlight', ...}],
            charts?: [{type: 'bar'|'line'|'pie'|'scatter'|'area', title?,
                       categories, series: [...], anchor?}],
        }

    Every sheet gets a bold, dark, centered header row, frozen so it stays
    visible while scrolling, and auto-sized columns — a plain, readable
    default rather than bare unstyled cells. Uses openpyxl's real Workbook
    API throughout, never writes raw bytes/CSV text into a .xlsx file;
    charts/tables/conditional formats are real, live Excel objects that
    stay correct if the underlying data is edited later, not static images.
    """
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)

    for sheet_spec in sheets:
        sheet_name = str(sheet_spec.get("name") or "Sheet")[:31]
        columns = sheet_spec.get("columns", [])
        rows = sheet_spec.get("rows", [])

        sheet = workbook.create_sheet(title=sheet_name)

        for col_index, header in enumerate(columns, start=1):
            cell = sheet.cell(row=1, column=col_index, value=header)
            cell.font = _HEADER_FONT
            cell.fill = _HEADER_FILL
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        for row_index, row in enumerate(rows, start=2):
            for col_index, value in enumerate(row, start=1):
                sheet.cell(row=row_index, column=col_index, value=value)

        for col_index, header in enumerate(columns, start=1):
            longest = max(
                [len(str(header))] + [len(str(row[col_index - 1])) for row in rows if col_index - 1 < len(row)],
                default=len(str(header)),
            )
            sheet.column_dimensions[get_column_letter(col_index)].width = min(max(longest + 2, 10), 45)

        if columns:
            sheet.freeze_panes = "A2"

        if sheet_spec.get("number_formats"):
            _apply_number_formats(sheet, columns, len(rows), sheet_spec["number_formats"])

        if sheet_spec.get("table"):
            _apply_table(sheet, columns, len(rows), sheet_spec["table"])

        if sheet_spec.get("conditional_formats"):
            _apply_conditional_formats(sheet, columns, len(rows), sheet_spec["conditional_formats"])

        if sheet_spec.get("charts"):
            _apply_charts(sheet, columns, len(rows), sheet_spec["charts"])

    workbook.save(str(path))
