from __future__ import annotations

import re
from pathlib import Path
from typing import Any

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH

_HEADING_RE = re.compile(r"^(#{1,3})\s+(.*)$")
_BULLET_RE = re.compile(r"^[-*]\s+(.*)$")
_NUMBERED_RE = re.compile(r"^\d+[.)]\s+(.*)$")
_BOLD_RE = re.compile(r"\*\*(.+?)\*\*")

_HEADING_LEVELS = {1: 1, 2: 2, 3: 3}


def _add_inline_runs(paragraph, text: str) -> None:
    """Split on '**bold**' spans and add each piece as its own run, same
    supported markdown subset as _pdf_writer's create_pdf.
    """
    pos = 0
    for match in _BOLD_RE.finditer(text):
        if match.start() > pos:
            paragraph.add_run(text[pos : match.start()])
        run = paragraph.add_run(match.group(1))
        run.bold = True
        pos = match.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])


def _write_markdown(document: docx.Document, markdown: str) -> None:
    for raw_line in markdown.split("\n"):
        heading_match = _HEADING_RE.match(raw_line)
        bullet_match = _BULLET_RE.match(raw_line)
        numbered_match = _NUMBERED_RE.match(raw_line)

        if heading_match:
            level = len(heading_match.group(1))
            document.add_heading(heading_match.group(2), level=_HEADING_LEVELS[level])
        elif bullet_match:
            p = document.add_paragraph(style="List Bullet")
            _add_inline_runs(p, bullet_match.group(1))
        elif numbered_match:
            p = document.add_paragraph(style="List Number")
            _add_inline_runs(p, numbered_match.group(1))
        elif raw_line.strip():
            p = document.add_paragraph()
            _add_inline_runs(p, raw_line)
        else:
            document.add_paragraph()


def create_docx(
    path: Path, title: str, content: str, tables: list[dict[str, Any]] | None = None
) -> None:
    """Build a real, valid Word document at `path` from a title, simple-markdown
    content, and optional tables — mirrors _pdf_writer.create_pdf's supported
    markdown subset ('#'/'##'/'###' headings, '**bold**', '-'/'*'/numbered
    lists) so callers can reuse the same content across both skills.

    Uses python-docx's Document API — never writes raw text/markdown bytes
    directly to a .docx file, which would produce a corrupt, non-Word file
    despite the extension.
    """
    document = docx.Document()

    title_paragraph = document.add_heading(title, level=0)
    title_paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT

    _write_markdown(document, content)

    for table_spec in tables or []:
        headers = table_spec.get("headers", [])
        rows = table_spec.get("rows", [])
        if table_spec.get("title"):
            document.add_heading(table_spec["title"], level=2)

        table = document.add_table(rows=1 if headers else 0, cols=len(headers) or (len(rows[0]) if rows else 1))
        table.style = "Light Grid Accent 1"

        if headers:
            for cell, header_text in zip(table.rows[0].cells, headers):
                cell.text = str(header_text)
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.bold = True

        for row in rows:
            cells = table.add_row().cells
            for cell, value in zip(cells, row):
                cell.text = str(value)

    document.save(str(path))
