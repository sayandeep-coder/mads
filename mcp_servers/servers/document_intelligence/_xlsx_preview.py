from __future__ import annotations

from pathlib import Path
from typing import Any

import openpyxl

# UI-facing, not tool-facing (contrast with _xlsx.py's read_xlsx, which is
# what the excel skill exposes to the model): reads every sheet at once so
# the web preview panel can offer a sheet switcher, and treats row 1 as a
# header row for display purposes, which read_xlsx deliberately doesn't
# (a tool has no business guessing what counts as a "header").
_MAX_ROWS_PER_SHEET = 500


def read_all_sheets_for_preview(path: Path) -> dict[str, Any]:
    workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)
    try:
        sheets: list[dict[str, Any]] = []
        for sheet_name in workbook.sheetnames:
            sheet = workbook[sheet_name]
            rows: list[list[Any]] = []
            for row_index, row in enumerate(sheet.iter_rows(values_only=True)):
                if row_index >= _MAX_ROWS_PER_SHEET:
                    break
                rows.append(["" if cell is None else cell for cell in row])

            headers = rows[0] if rows else []
            data_rows = rows[1:] if rows else []

            sheets.append(
                {
                    "name": sheet_name,
                    "headers": [str(h) for h in headers],
                    "rows": data_rows,
                    "total_rows": sheet.max_row,
                    "truncated": sheet.max_row > _MAX_ROWS_PER_SHEET,
                }
            )

        return {"sheets": sheets}
    finally:
        workbook.close()
