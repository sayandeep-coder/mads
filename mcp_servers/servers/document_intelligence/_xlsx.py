from __future__ import annotations

from pathlib import Path
from typing import Any

import openpyxl

_MAX_ROWS = 500


def read_xlsx(
    path: Path,
    sheet_name: str | None = None,
    start_row: int | None = None,
    end_row: int | None = None,
) -> dict[str, Any]:
    """Read rows from a spreadsheet. Defaults to the first sheet and the first _MAX_ROWS rows."""
    workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)
    try:
        sheet = workbook[sheet_name] if sheet_name else workbook[workbook.sheetnames[0]]

        first = max((start_row or 1) - 1, 0)
        last = min(end_row or sheet.max_row, sheet.max_row)
        truncated = (last - first) > _MAX_ROWS
        if truncated:
            last = first + _MAX_ROWS

        rows: list[list[Any]] = []
        for row_index, row in enumerate(sheet.iter_rows(values_only=True)):
            if row_index < first:
                continue
            if row_index >= last:
                break
            rows.append(list(row))

        result: dict[str, Any] = {
            "sheet_name": sheet.title,
            "all_sheet_names": workbook.sheetnames,
            "total_rows": sheet.max_row,
            "total_columns": sheet.max_column,
            "rows_read": f"{first + 1}-{last}",
            "rows": rows,
            "truncated": truncated,
        }
        if truncated:
            result["note"] = f"Row list truncated at {_MAX_ROWS} rows. Specify start_row/end_row to read a narrower range."
        return result
    finally:
        workbook.close()
