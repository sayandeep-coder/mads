from __future__ import annotations

import re

from pptx.util import Pt

from mcp_servers.servers.document_intelligence import _pptx_theme as theme

_HEADING_RE = re.compile(r"^(#{1,2})\s+(.*)$")
_BULLET_RE = re.compile(r"^[-*]\s+(.*)$")
_BOLD_RE = re.compile(r"\*\*(.+?)\*\*")


def apply_bold_runs(paragraph, text: str, base_size, base_color, bold: bool = False) -> None:
    """Split text on '**bold**' spans and add them as separate runs with real bold
    formatting, instead of leaving literal asterisks in the rendered text.
    """
    text = str(text)
    cursor = 0
    any_run_added = False

    for match in _BOLD_RE.finditer(text):
        if match.start() > cursor:
            run = paragraph.add_run()
            run.text = text[cursor : match.start()]
            run.font.size = base_size
            run.font.color.rgb = base_color
            run.font.name = theme.FONT_FAMILY
            run.font.bold = bold
            any_run_added = True

        run = paragraph.add_run()
        run.text = match.group(1)
        run.font.size = base_size
        run.font.color.rgb = base_color
        run.font.name = theme.FONT_FAMILY
        run.font.bold = True
        any_run_added = True
        cursor = match.end()

    if cursor < len(text) or not any_run_added:
        run = paragraph.add_run()
        run.text = text[cursor:]
        run.font.size = base_size
        run.font.color.rgb = base_color
        run.font.name = theme.FONT_FAMILY
        run.font.bold = bold


def add_markdown_bullets(text_frame, lines: list[str], base_size, base_color) -> None:
    """Render a list of bullet lines into a text frame, with real formatting:

    - A line starting with '#'/'##' becomes a bold section heading at level 0/1,
      acting as a group label for the plain bullets that follow.
    - A line starting with '-'/'*' (or any plain line) becomes a normal bullet,
      indented one level below the most recent heading.
    - '**bold**' spans anywhere are rendered as real bold runs.

    This is the bullets-slide and comparison-slide equivalent of the
    markdown renderer already used for Google Docs and generated PDFs.
    """
    raw_lines = [line for item in lines for line in str(item).splitlines()]
    first_paragraph = True
    heading_level: int | None = None

    for raw_line in raw_lines:
        if not raw_line.strip():
            continue
        heading_match = _HEADING_RE.match(raw_line)
        bullet_match = _BULLET_RE.match(raw_line)
        content = heading_match.group(2) if heading_match else (bullet_match.group(1) if bullet_match else raw_line)

        paragraph = text_frame.paragraphs[0] if first_paragraph else text_frame.add_paragraph()
        first_paragraph = False
        paragraph.space_after = Pt(14) if heading_match else Pt(10)

        if heading_match:
            heading_level = len(heading_match.group(1)) - 1
            paragraph.level = heading_level
            apply_bold_runs(paragraph, content, base_size, base_color, bold=True)
        else:
            paragraph.level = heading_level + 1 if heading_level is not None else 0
            prefix_run = paragraph.add_run()
            prefix_run.text = "•  "
            prefix_run.font.size = base_size
            prefix_run.font.color.rgb = base_color
            prefix_run.font.name = theme.FONT_FAMILY
            apply_bold_runs(paragraph, content, base_size, base_color)
